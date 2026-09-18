# -*- coding: utf-8 -*-
"""
tests/test_catalog_age.py — возраст данных в Каталоге по ПРАВИЛАМ сбора.

Job «Listing Suite Auto Collector» (13:00 Kyiv) работает верно, врал
экран. Проверено по его ноутбуку 17.09: пары с продажами за 30 дней
(`collection_tier = 'daily'`) собираются каждый день, остальные —
раз в неделю в день, закреплённый хэшем (`weekly_day`, 0 = понедельник);
`status = 'wound_down'` не собирается вовсе. Экран этого не знал и
показывал свёрнутый товар как отставший на три недели, а новую пару,
ждущую свой четверг, — как «не собиралась».

Проверяются четыре состояния карточки и полоса по рынку:
  · свёрнутый — словом, без возраста и без попадания в «отстают»;
  · ежедневный старше двух дней — отстаёт, недельный в те же два — нет;
  · собранный — дата, возраст словами и график («по четвергам»);
  · новый — «первый сбор в четверг», а не «не собирался».

Запуск (pytest не нужен):  python tests/test_catalog_age.py
"""
from __future__ import annotations

import json
import pathlib
import sys

import pandas as pd

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import services.db                                     # noqa: E402
services.db.get_conn = lambda: type(
    "C", (), {"close": lambda self: None})()

FAILS: list[str] = []


def check(name: str, cond: bool) -> None:
    print(("  OK   " if cond else "  FAIL ") + name)
    if not cond:
        FAILS.append(name)


RAW = json.dumps({"images": ["a"], "number_of_videos": 1, "aplus": True,
                  "average_rating": "4,5", "price": "10", "sold_by": "Dnipro-M"})
NOW = pd.Timestamp.now("UTC")


def product(asin, mp, days, status="active", tier="weekly", day=3):
    ts = pd.NaT if days is None else (NOW - pd.Timedelta(days=days, hours=1))
    return dict(sku_group=f"175{asin}", asin=asin, marketplace=mp,
                is_competitor=False, status=status, collection_tier=tier,
                weekly_day=day, fetched_at=ts, ok=True if days is not None else None,
                title="T" if days is not None else None, in_stock=True,
                review_count=100, is_amazon_choice=False,
                raw=RAW if days is not None else None)


# ES: сегодня; недельная 4 дня (норма); недельная 10 дней (отстаёт);
# ежедневная 3 дня (отстаёт — для неё это много); новая ждёт четверга;
# свёрнутая 20 дней (не отстаёт — её не собирают). IT: всё свежее.
CAT = pd.DataFrame([
    product("B0TODAY", "es", 0, tier="daily"),
    product("B0FOUR", "es", 4, day=3),
    product("B0TEN", "es", 10, day=1),
    product("B0DAILY3", "es", 3, tier="daily"),
    product("B0NEVER", "es", None, day=3),
    product("B0WOUND", "es", 20, status="wound_down"),
    product("B0IT1", "it", 1, day=0),
    product("B0IT2", "it", 2, tier="daily"),
])


def fake_sql(sql, conn, **kw):
    s = str(sql)
    if "FROM product_matrix m" in s and "is_amazon_choice" in s:
        return CAT.copy()
    if "count(*) AS pairs" in s:                    # план сбора по рынкам
        mps = list((kw.get("params") or {}).get("mps") or [])
        act = CAT[(CAT["status"] == "active") & CAT["marketplace"].isin(mps)]
        return act.groupby("marketplace").size().rename("pairs").reset_index()
    return pd.DataFrame()


pd.read_sql = fake_sql
from streamlit.testing.v1 import AppTest              # noqa: E402

at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=180).run()
at.switch_page("pages/catalog.py").run()
check("страница отрисована без ошибки", not at.exception)

html = " ".join(str(m.value) for m in at.markdown)

# --- 1. полоса по рынкам
strip = next((str(m.value) for m in at.markdown
              if "<b style=" in str(m.value) and ">ES</b>" in str(m.value)), "")
check("полоса возраста по рынкам есть", bool(strip) and ">IT</b>" in strip)
es = strip.split(">ES</b>")[1].split(">IT</b>")[0]
check("ES: отстают от графика двое — недельная 10 дн. и ежедневная 3 дн.",
      "отстают от графика: 2" in es)
check("ES: свёрнутая не в «отстают», а названа отдельно",
      "свёрнутых: 1" in es and "до 10 дн." in es)
check("ES: новая — «ждут первого сбора», не «не собирались»",
      "ждут первого сбора: 1" in es and "не собирались" not in es)
it = strip.split(">IT</b>")[1]
check("IT: без отставаний — ежедневная в 2 дня ещё в норме", "отстают" not in it)
check("и сказано, по какому правилу идёт сбор",
      any("раз в неделю" in str(c.value) and "свёрнутые" in str(c.value) for c in at.caption))


def card(asin: str) -> str:
    return next((str(m.value) for m in at.markdown
                 if asin in str(m.value) and ("назад" in str(m.value) or "сегодня" in str(m.value)
                                              or "первый сбор" in str(m.value) or "свёрнут" in str(m.value))), "")


# --- 2. подпись на карточке
src = (ROOT / "pages/catalog.py").read_text(encoding="utf-8")
warn = src.split('WARN_TEXT = "')[1].split('"')[0]
muted = src.split('MUTED = "')[1].split('"')[0]
check("сегодняшняя ежедневная — «сегодня · обновляется ежедневно»",
      "сегодня" in card("B0TODAY") and "обновляется ежедневно" in card("B0TODAY"))
check("недельная 4 дня — «4 дня назад · по четвергам», не отстаёт",
      "4 дня назад" in card("B0FOUR") and "по четвергам" in card("B0FOUR")
      and f'color:{muted};">4 дня назад' in card("B0FOUR"))
check("недельная 10 дней — отстаёт, выделена",
      f'color:{warn};">10 дней назад' in card("B0TEN") and "по вторникам" in card("B0TEN"))
check("ежедневная 3 дня — отстаёт, хотя по календарю это немного",
      f'color:{warn};">3 дня назад' in card("B0DAILY3"))
check("новая — «первый сбор в четверг», не «не собирался»",
      "первый сбор в четверг" in card("B0NEVER") and "не собирался" not in card("B0NEVER"))
c = card("B0WOUND")
check("свёрнутая — словом, без возраста",
      "свёрнут" in c and "назад" not in c and "последний сбор" in c)
check("и приглушённым цветом, не предупреждающим",
      f'color:{muted};">свёрнут' in c)

# --- 3. выгрузка — ровно то, что на экране, в CSV и Excel
# Число в подписи кнопок обязано совпадать с числом строк под фильтром:
# файл «всего каталога» вместо отфильтрованного — тихая подмена.
# Байты из download_button в AppTest не достать, поэтому сборщик XLSX
# проверяется отдельно на том же кадре: файл открывается и в нём
# колонки сбора, которые подписаны на карточках.
_dl = {b.key: str(b.label) for b in at.get("download_button")}
check("две кнопки выгрузки: CSV и Excel",
      "cat-export-csv" in _dl and "cat-export-xlsx" in _dl)
check(f"без фильтра — все 8 строк ({_dl.get('cat-export-csv')})",
      _dl["cat-export-csv"].endswith("· 8") and _dl["cat-export-xlsx"].endswith("· 8"))
at.multiselect[0].set_value(["it"]).run()
_dl = {b.key: str(b.label) for b in at.get("download_button")}
check(f"с фильтром IT — две строки в обоих ({_dl.get('cat-export-xlsx')})",
      _dl["cat-export-csv"].endswith("· 2") and _dl["cat-export-xlsx"].endswith("· 2"))
# Отметил рынок для сбора — выгрузка идёт ПО НЕМУ, без промежуточного
# нажатия: «выбрал PL, собрал, скачал по нему же». Список при этом
# не режется (собирать три рынка и смотреть весь каталог — нормально),
# а подпись кнопки говорит, что в файле: «CSV · 2 (IT)».
at.multiselect[0].set_value([]).run()
next(c for c in at.checkbox if c.key == "collect-mp-it").set_value(True).run()
# Нажимая «собрать», надо видеть, когда собирали в прошлый раз: рядом
# с числом — дата самого свежего снапшота рынка (у IT — вчера).
_plan = next((str(m.value) for m in at.markdown
              if "white-space" in str(m.value) and "товар" in str(m.value)), "")
_it_last = (NOW - pd.Timedelta(days=1, hours=1)).strftime("%d.%m")
check(f"рядом с числом — дата последнего сбора рынка ({_plan[-60:]})",
      "последний сбор" in _plan and _it_last in _plan)

_dl = {b.key: str(b.label) for b in at.get("download_button")}
check(f"галочка IT сужает выгрузку до IT и говорит об этом ({_dl.get('cat-export-csv')})",
      _dl["cat-export-csv"].endswith("· 2 (IT)") and _dl["cat-export-xlsx"].endswith("· 2 (IT)"))
check("список при этом не режется — 8 строк на экране",
      any("8 товаров" in str(m.value) for m in at.markdown))
check("и подсказка называет рынки, а не «то, что на экране»",
      all("IT" in str(b.help) and "то, что на экране" not in str(b.help)
          for b in at.get("download_button")))
# ссылка «показать в списке» остаётся для просмотра — и режет уже список
next(b for b in at.button if b.key == "collect-show").click().run()
check("«показать эти рынки в списке» ставит фильтр списка на IT",
      at.multiselect[0].value == ["it"])
next(c for c in at.checkbox if c.key == "collect-mp-it").set_value(False).run()
_dl = {b.key: str(b.label) for b in at.get("download_button")}
check(f"сняли галочку — выгрузка снова по фильтрам списка ({_dl.get('cat-export-csv')})",
      _dl["cat-export-csv"].endswith("· 2") and "(IT)" not in _dl["cat-export-csv"])
# кнопки стоят в ряду со сбором, пояснение — в подсказке кнопки
# в табличном виде дата сбора — колонкой, как на карточке
at.session_state["cat_mode"] = "table"
at.run()
_tbl = at.dataframe[0].value if at.dataframe else None
check("в таблице есть колонка «Собрано» с датой снапшота",
      _tbl is not None and "fetched" in _tbl.columns
      and _tbl["fetched"].dropna().astype(str).str.match(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}").all())
at.session_state["cat_mode"] = "cards"
at.run()
check("и сказано, что выгружается то, что на экране",
      all("то, что на экране" in str(b.help) for b in at.get("download_button")))
check("выгрузка стоит в ряду со сбором и выделена",
      all(b.proto.type == "primary" for b in at.get("download_button")))

import io, re  # noqa: E402
_src = (ROOT / "pages/catalog.py").read_text(encoding="utf-8")
_fn = "def _xlsx_bytes" + _src.split("def _xlsx_bytes")[1].split("\n\n\n")[0]
from services.marketplaces import product_url  # noqa: E402
_ns = {"pd": pd, "product_url": product_url}; exec(_fn, _ns)
_frame = pd.DataFrame([{"asin": "B0X", "status": "wound_down", "tier": "weekly",
                        "weekly_day": 3, "age_days": 20, "name": "Тест"}])
_wb = pd.read_excel(io.BytesIO(_ns["_xlsx_bytes"](_frame)), sheet_name="catalog")
check("XLSX открывается и держит колонки сбора и кириллицу",
      list(_wb.columns) == ["asin", "status", "tier", "weekly_day", "age_days", "name"]
      and _wb.iloc[0]["name"] == "Тест" and int(_wb.iloc[0]["weekly_day"]) == 3)
_frame = pd.DataFrame([{"asin": "B0G4S9SJ3M", "mp": "be", "name": "x"},
                       {"asin": "B0G4S9SJ3M", "mp": "es", "name": "y"}])
import openpyxl  # noqa: E402
_ws = openpyxl.load_workbook(io.BytesIO(_ns["_xlsx_bytes"](_frame)))["catalog"]
_links = [(_ws.cell(row=r, column=1).value, getattr(_ws.cell(row=r, column=1).hyperlink, "target", None))
          for r in (2, 3)]
check(f"ASIN в XLSX — гиперссылка, в ячейке сам ASIN ({_links[1]})",
      _links[1] == ("B0G4S9SJ3M", "https://www.amazon.es/dp/B0G4S9SJ3M"))
# Бельгия — витрина amazon.com.be: шаблон amazon.{mp} дал бы amazon.be
check(f"Бельгия ведёт на amazon.com.be, не amazon.be ({_links[0][1]})",
      _links[0][1] == "https://www.amazon.com.be/dp/B0G4S9SJ3M")
check("в CSV ссылок нет — ASIN как текст",
      "amazon." not in _src.split("exp.to_csv")[0][-400:] and 'exp.to_csv(index=False)' in _src)
check("в выгрузке есть колонки сбора",
      all(f'"{c}":' in _src.split("exp = pd.DataFrame")[1].split("] for x in rows")[0]
          for c in ("status", "tier", "weekly_day", "fetched_at", "age_days")))

print()
print("ИТОГ:", "все проверки прошли" if not FAILS
      else f"{len(FAILS)} провалов: {FAILS}")
sys.exit(1 if FAILS else 0)
